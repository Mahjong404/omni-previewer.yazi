import datetime
import io
import os
import sys
import time
import zipfile
import unicodedata
from xml.etree import ElementTree
from openpyxl import load_workbook
from openpyxl.styles.colors import COLOR_INDEX

MAX_SHEETS, MAX_ROWS, MAX_COLS, CELL_MAX = 20, 200, 12, 24
DIM, RESET, BOLD, TITLE = "\x1b[90m", "\x1b[0m", "\x1b[1m", "\x1b[1;36m"


def width(text):
    return sum(2 if unicodedata.east_asian_width(c) in "WF" else 1 for c in text)


def fit(text, w):
    if width(text) <= w:
        return text + " " * (w - width(text))
    out, n = "", 0
    for c in text:
        cw = 2 if unicodedata.east_asian_width(c) in "WF" else 1
        if n + cw > w - 1:
            break
        out += c
        n += cw
    return out + "…" + " " * (w - n - 1)


def fit_center(text, w):
    t = fit(text, w).rstrip()
    lead = max(0, (w - width(t)) // 2)
    return " " * lead + t + " " * max(0, w - lead - width(t))


def theme_palette(path):
    try:
        with zipfile.ZipFile(path) as z:
            root = ElementTree.fromstring(z.read("xl/theme/theme1.xml"))
    except Exception:
        return []
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    scheme = root.find(".//a:clrScheme", ns)
    names = ["dk1", "lt1", "dk2", "lt2", "accent1", "accent2", "accent3", "accent4",
             "accent5", "accent6", "hlink", "folHlink"]
    colors = []
    for name in names:
        el = scheme.find("a:" + name, ns) if scheme is not None else None
        val = None
        if el is not None:
            srgb, sysc = el.find("a:srgbClr", ns), el.find("a:sysClr", ns)
            if srgb is not None:
                val = srgb.get("val")
            elif sysc is not None:
                val = sysc.get("lastClr") or {"windowText": "000000", "window": "FFFFFF"}.get(sysc.get("val"))
        colors.append(val)
    if len(colors) >= 4:
        colors[0], colors[1] = colors[1], colors[0]
        colors[2], colors[3] = colors[3], colors[2]
    return colors


def resolve(color, palette):
    if color is None:
        return None
    raw = None
    if color.type == "rgb" and isinstance(color.rgb, str):
        raw = color.rgb[-6:]
    elif color.type == "theme" and 0 <= color.theme < len(palette) and palette[color.theme]:
        raw = palette[color.theme]
    elif color.type == "indexed" and 0 <= color.indexed < len(COLOR_INDEX):
        raw = COLOR_INDEX[color.indexed][-6:]
    if not raw:
        return None
    try:
        rgb = tuple(int(raw[i:i + 2], 16) for i in (0, 2, 4))
    except ValueError:
        return None
    tint = getattr(color, "tint", 0) or 0
    if tint:
        rgb = tuple(round(min(255, max(0, c * (1 - tint) + 255 * tint if tint > 0 else c * (1 + tint)))) for c in rgb)
    return rgb


def style_seq(cell, palette):
    seq = ""
    fill = cell.fill
    if fill is not None and fill.patternType and fill.patternType != "none":
        rgb = resolve(fill.fgColor, palette)
        if rgb and rgb != (255, 255, 255):
            seq += f"\x1b[48;2;{rgb[0]};{rgb[1]};{rgb[2]}m"
    font = cell.font
    if font is not None:
        rgb = resolve(font.color, palette)
        if rgb and rgb != (0, 0, 0):
            seq += f"\x1b[38;2;{rgb[0]};{rgb[1]};{rgb[2]}m"
        if font.bold:
            seq += BOLD
    return seq


def text_of(cell):
    v = cell.value
    if v is None:
        return ""
    if isinstance(v, datetime.datetime):
        return v.strftime("%Y-%m-%d") if v.time() == datetime.time(0, 0) else v.strftime("%Y-%m-%d %H:%M")
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    return " ".join(str(v).split())


def render_sheet(sheet, palette):
    ncols = min(sheet.max_column or 0, MAX_COLS)
    nrows = min(sheet.max_row or 0, MAX_ROWS)
    if not ncols or not nrows:
        print("(empty sheet)")
        return
    rows = list(sheet.iter_rows(min_row=1, max_row=nrows, max_col=ncols))
    grid = [[text_of(c) for c in row] for row in rows]
    while grid and not any(grid[-1]):
        grid.pop()
        rows.pop()
    last = 0
    for r in grid:
        for i, v in enumerate(r):
            if v:
                last = i + 1
    ncols = min(ncols, last)
    nrows = len(grid)
    if not ncols or not nrows:
        print("(empty sheet)")
        return

    anchor_span, covered, centered = {}, {}, set()
    try:
        ranges = sheet.merged_cells.ranges
    except Exception:
        ranges = []
    for rng in ranges:
        r0, c0 = rng.min_row, rng.min_col
        r1, c1 = min(rng.max_row, nrows), min(rng.max_col, ncols)
        if r0 > nrows or c0 > ncols:
            continue
        anchor = rows[r0 - 1][c0 - 1]
        anchor_span[(r0, c0)] = c1 - c0 + 1
        if getattr(anchor.alignment, "horizontal", None) in ("center", "centerContinuous"):
            centered.add((r0, c0))
        for rr in range(r0, r1 + 1):
            for cc in range(c0, c1 + 1):
                if (rr, cc) != (r0, c0):
                    covered[(rr, cc)] = anchor

    letters = [chr(65 + i) for i in range(ncols)]
    num_w = max(2, len(str(nrows)))
    colw = []
    for i in range(ncols):
        w = width(letters[i])
        for r in range(nrows):
            if (r + 1, i + 1) not in covered and anchor_span.get((r + 1, i + 1), 1) == 1:
                w = max(w, width(grid[r][i]))
        colw.append(min(w, CELL_MAX))
    widths = [num_w] + colw

    def border(left, mid, right):
        print(DIM + left + mid.join("─" * (w + 2) for w in widths) + right + RESET)

    def emit(text, st, w):
        return " " + (st + text + RESET if st else text) + " " + DIM + "│" + RESET

    def line(r, header=False):
        out = DIM + "│" + RESET
        if header:
            out += emit(" " * num_w, "", num_w)
            for i in range(ncols):
                out += emit(fit(letters[i], colw[i]), BOLD, colw[i])
            print(out)
            return
        out += emit(f"{r:>{num_w}}", DIM, num_w)
        c = 1
        while c <= ncols:
            anchor = covered.get((r, c))
            if anchor is not None:
                out += emit(" " * colw[c - 1], style_seq(anchor, palette), colw[c - 1])
                c += 1
                continue
            span = min(anchor_span.get((r, c), 1), ncols - c + 1)
            cell = rows[r - 1][c - 1]
            st = style_seq(cell, palette)
            if span > 1:
                w = sum(colw[c - 1:c - 1 + span]) + 3 * (span - 1)
                text = fit_center(grid[r - 1][c - 1], w) if (r, c) in centered else fit(grid[r - 1][c - 1], w)
                out += emit(text, st, w)
                c += span
            else:
                out += emit(fit(grid[r - 1][c - 1], colw[c - 1]), st, colw[c - 1])
                c += 1
        print(out)

    border("┌", "┬", "┐")
    line(0, header=True)
    border("├", "┼", "┤")
    for r in range(1, nrows + 1):
        line(r)
    border("└", "┴", "┘")
    if (sheet.max_row or 0) > nrows or (sheet.max_column or 0) > ncols:
        print(DIM + f"[showing first {nrows}x{ncols} of {sheet.max_row}x{sheet.max_column}]" + RESET)


def ensure_xlsx(path, cache):
    """Legacy .xls (CFB binary) can't be read by openpyxl — convert it once via
    Excel COM to a cached .xlsx, then reuse the normal pipeline."""
    if not str(path).lower().endswith(".xls"):
        return path
    target = ((cache[:-5] if cache and cache.endswith(".ansi") else cache) or path + ".conv") + ".xlsx"
    if os.path.exists(target):
        return target
    import pythoncom
    import win32com.client

    pythoncom.CoInitialize()
    xl = win32com.client.DispatchEx("Excel.Application")
    try:
        xl.DisplayAlerts = 0
        xl.Visible = False
        wb = xl.Workbooks.Open(str(path), ReadOnly=True)
        tmp = target + "." + str(os.getpid()) + ".tmp.xlsx"
        os.makedirs(os.path.dirname(target), exist_ok=True)
        wb.SaveAs(tmp, 51)  # xlOpenXMLWorkbook (.xlsx)
        wb.Close(False)  # SaveAs swaps the handle to tmp; close before moving it
        os.replace(tmp, target)
    finally:
        try:
            xl.Quit()
        finally:
            pythoncom.CoUninitialize()
    return target


def prune_cache(d):
    try:
        files = [os.path.join(d, f) for f in os.listdir(d) if f.endswith(".ansi")]
        now = time.time()
        alive = []
        for f in files:
            try:
                if now - os.path.getmtime(f) > 7 * 86400:
                    os.remove(f)
                else:
                    alive.append((os.path.getmtime(f), f))
            except OSError:
                pass
        for _, f in sorted(alive, reverse=True)[100:]:
            try:
                os.remove(f)
            except OSError:
                pass
    except Exception:
        pass


buf = io.StringIO()
try:
    real = sys.stdout
    sys.stdout = buf
    try:
        src = ensure_xlsx(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else None)
        palette = theme_palette(src)
        book = load_workbook(src, data_only=False, keep_links=False)
        try:
            for i, sheet in enumerate(book.worksheets):
                if i >= MAX_SHEETS:
                    print(DIM + "[additional sheets omitted]" + RESET)
                    break
                if i:
                    print()
                print(TITLE + sheet.title + RESET)
                render_sheet(sheet, palette)
        finally:
            book.close()
    finally:
        sys.stdout = real
    out = buf.getvalue()
    if len(sys.argv) > 2:
        cache = sys.argv[2]
        try:
            d = os.path.dirname(cache)
            os.makedirs(d, exist_ok=True)
            tmp = cache + "." + str(os.getpid()) + ".tmp"
            with open(tmp, "w", encoding="utf-8", newline="") as f:
                f.write(out)
            os.replace(tmp, cache)
            prune_cache(d)
        except Exception:
            pass
    sys.stdout.write(out)
except Exception as exc:
    sys.stdout = sys.__stdout__
    print("XLSX preview failed: " + str(exc), file=sys.stderr)
    sys.exit(1)
